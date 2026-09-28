from pathlib import Path
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path("exports/Lexy_Learning_Growth_Manual_Test_Cases.xlsx")

HEADERS = [
    "Test Case ID", "Stage", "Module", "Test Case", "Objective", "Preconditions",
    "Test Data", "Steps", "Expected Result", "Priority", "Test Type", "Role",
    "Platform", "Automation Reference", "Status", "Actual Result", "Tester",
    "Execution Date", "Defect ID", "Notes",
]

cases = []
sequence = {}


def add(stage, module, title, objective, preconditions, data, steps, expected,
        priority="High", test_type="Functional", role="Candidate",
        platform="Web - Desktop & Mobile", automation=""):
    sequence[stage] = sequence.get(stage, 0) + 1
    prefixes = {
        "Programme": "LGP", "Stage A": "LGA", "Stage B": "LGB",
        "Stage C": "LGC", "Stage D": "LGD", "Rewards": "LGR",
        "Privacy & Security": "LGS", "Resilience": "LGX",
    }
    cases.append({
        "Test Case ID": f"{prefixes[stage]}-{sequence[stage]:03d}",
        "Stage": stage, "Module": module, "Test Case": title,
        "Objective": objective, "Preconditions": preconditions, "Test Data": data,
        "Steps": steps, "Expected Result": expected, "Priority": priority,
        "Test Type": test_type, "Role": role, "Platform": platform,
        "Automation Reference": automation, "Status": "Not Run",
        "Actual Result": "", "Tester": "", "Execution Date": "",
        "Defect ID": "", "Notes": "",
    })


def numbered(*items):
    return "\n".join(f"{i}. {item}" for i, item in enumerate(items, 1))


# Programme access, rollout, navigation, and compatibility
add("Programme", "Access", "Unauthenticated visitor is redirected to sign-in",
    "Confirm private learning pages cannot be opened anonymously.", "Logged out.",
    "Direct URL to Learning & Growth.",
    numbered("Open the Learning & Growth URL in a private browser window.",
             "Observe navigation and network responses."),
    "No learning data is rendered. User is taken to sign-in or receives the approved unauthenticated state. No private response is cached.",
    "Critical", "Security", "Anonymous", automation="learning-growth.test.ts")
for role in ["Recruiter", "Tenant admin", "Platform admin"]:
    add("Programme", "Role access", f"{role} cannot open candidate-private learning",
        "Ensure staff roles cannot inspect private candidate development records.",
        f"Signed in as {role}; candidate learning data exists.", "Candidate with completed learning.",
        numbered("Sign in with the stated staff role.", "Open the candidate Learning & Growth URL.",
                 "Try direct URLs for courses, assessments, voice progress, and achievements."),
        "Access is denied without revealing whether private learning records exist.",
        "Critical", "Security", role, automation="learning-growth.test.ts")
add("Programme", "Pilot rollout", "Programme disabled globally",
    "Verify the feature flag closes every learning entry point.", "Pilot flag disabled.",
    "Eligible candidate in allowlist.",
    numbered("Sign in as the eligible candidate.", "Open programme and deep links.",
             "Attempt to save interests or enrol."),
    "Programme shows unavailable; no course/assessment/voice data is exposed; writes do not create records.",
    "Critical", "Configuration", automation="learning-growth.test.ts")
add("Programme", "Pilot rollout", "Candidate not in pilot allowlist",
    "Verify per-candidate rollout is fail-closed.", "Pilot enabled; candidate omitted from allowlist.",
    "Candidate with otherwise valid prerequisites.",
    numbered("Sign in as omitted candidate.", "Open programme and all deep links."),
    "Programme is unavailable and no private learning records are returned or changed.",
    "Critical", "Security", automation="learning-growth.test.ts")
add("Programme", "Pilot rollout", "Wildcard allowlist admits eligible candidates",
    "Confirm the wildcard rollout enables the intended population.", "Pilot enabled; allowlist is '*'.",
    "Two eligible candidates in different tenants.",
    numbered("Sign in as each candidate.", "Open Learning & Growth."),
    "Each candidate can open only their own programme and receives private, no-store responses.",
    "High", "Configuration")
add("Programme", "Navigation", "Refresh preserves current programme stage",
    "Verify durable progress survives reload.", "Candidate has partially completed the programme.",
    "Saved interests, draft lesson, draft assessment, or active voice cycle.",
    numbered("Open the relevant stage.", "Refresh the browser.", "Sign out and sign back in."),
    "The correct stage, progress, and resume point are restored without duplicating work.",
    "High", "Lifecycle")
add("Programme", "Navigation", "Browser back and forward do not duplicate writes",
    "Prevent accidental duplicate enrolments, attempts, or rewards.", "Candidate can perform a learning mutation.",
    "Any enrol/start/complete action.",
    numbered("Complete the action.", "Use Back, then Forward.", "Refresh the resulting page."),
    "Only one durable record and one applicable reward exist; the UI reflects the existing result.",
    "High", "Lifecycle")
add("Programme", "Caching", "Private responses are not browser/proxy cached",
    "Ensure one candidate cannot receive another candidate's cached data.", "Two candidate accounts available.",
    "Candidate A and Candidate B in the same tenant.",
    numbered("Open programme as Candidate A.", "Sign out.", "Sign in as Candidate B in the same browser.",
             "Inspect pages and response Cache-Control headers."),
    "Candidate B sees only their own data. Private endpoints return Cache-Control: private, no-store.",
    "Critical", "Security")
for viewport in ["390x844 mobile", "768x1024 tablet", "1440x900 desktop"]:
    add("Programme", "Responsive UI", f"Programme shell works at {viewport}",
        "Check primary navigation and content remain usable.", "Eligible candidate.",
        viewport, numbered("Set browser viewport.", "Open every programme stage available.",
                           "Scroll, use menus, and submit a form."),
        "No clipped controls, horizontal overflow, inaccessible dialogs, or hidden validation messages.",
        "Medium", "Usability", platform=viewport)

# Stage A: onboarding, interests, baseline gate, goals, plan
add("Stage A", "Interests", "First visit starts at interests",
    "Confirm the initial state is clear and empty.", "Eligible candidate; no learning profile.",
    "New candidate.",
    numbered("Open Learning & Growth.", "Review the stage indicator and fields."),
    "Stage A interests is active, revision is zero, no completion timestamp is fabricated, and suggestions are clearly marked.",
    automation="learning-growth.test.ts")
interest_validations = [
    ("Unknown request field", "Add an unsupported field such as baselineInterviewCompleted=true.",
     "Request is rejected; baseline is not bypassed; no profile is created or changed."),
    ("Fractional revision", "Submit revision 0.5.", "Validation rejects the request and preserves prior state."),
    ("Blank custom interest", "Select Other and enter spaces only.", "Inline validation requires meaningful text."),
    ("Blank immediate role", "Enter a role containing only spaces.", "Blank role is rejected and not persisted."),
    ("Overlong custom interest", "Enter text longer than the documented limit.", "Validation rejects excess length without truncating silently."),
    ("Unsupported career area", "Modify the request to include an unknown career-area value.", "Request is rejected; no mutation occurs."),
    ("Unsupported education stage", "Modify the request to include an unknown education-stage value.", "Request is rejected; no mutation occurs."),
    ("Graduation year below range", "Enter a year earlier than the accepted range.", "Field-level validation prevents save."),
    ("Graduation year above range", "Enter a year later than the accepted range.", "Field-level validation prevents save."),
    ("Duplicate immediate roles", "Enter the same role twice with different surrounding whitespace.",
     "Roles are trimmed and deduplicated; one canonical value is saved."),
]
for title, action, expected in interest_validations:
    add("Stage A", "Interests validation", title,
        "Verify strict input validation and atomic writes.", "Candidate is at interests with a known revision.",
        action, numbered("Record the current profile and revision.", action, "Submit.", "Reload the page."),
        expected + " The revision changes only after a valid save.", "High", "Negative")
add("Stage A", "Interests", "Save complete interest profile",
    "Confirm normalization and progression to baseline.", "Candidate is at interests.",
    "Customer service; graduating; communication priority; English; start now.",
    numbered("Complete every field with valid values.", "Save.", "Reload."),
    "Normalized values persist, revision increments exactly once, and the next stage is baseline.",
    "Critical", "Positive", automation="learning-growth.test.ts")
add("Stage A", "Concurrency", "Two tabs save the same interest revision",
    "Confirm optimistic concurrency prevents lost updates.", "Open the same candidate profile in two tabs.",
    "Tab A and Tab B both hold revision N.",
    numbered("Change interests differently in each tab.", "Save Tab A.", "Save Tab B without refreshing.",
             "Reload both tabs."),
    "Exactly one write succeeds. The stale tab receives a conflict/reload message and cannot overwrite the winner.",
    "Critical", "Concurrency")
add("Stage A", "Baseline gate", "Goals cannot be confirmed before baseline",
    "Ensure saved interests cannot substitute for authoritative baseline completion.",
    "Interests saved; baseline not completed.", "Valid immediate, 3-year, and 5-year goals.",
    numbered("Deep-link to goals or submit goals through the API/UI.", "Reload programme."),
    "Confirmation is blocked; goals are not marked confirmed and no plan is generated.",
    "Critical", "Lifecycle", automation="learning-growth.test.ts")
add("Stage A", "Baseline gate", "Baseline completion unlocks goals",
    "Confirm authoritative baseline completion advances the journey.", "Interests saved.",
    "Complete the real baseline interview flow.",
    numbered("Complete baseline.", "Return to Learning & Growth.", "Refresh."),
    "Baseline is shown completed using the real completion state/time, and goals can be confirmed.",
    "Critical", "Positive")
goal_validations = [
    ("Blank immediate goal", "Immediate goal is spaces only."),
    ("Blank three-year goal", "Three-year goal is spaces only."),
    ("Blank five-year goal", "Five-year goal is spaces only."),
    ("Unknown goal field", "Include baselineInterviewCompleted or another unsupported property."),
    ("Stale goal revision", "Submit an older revision after another valid save."),
]
for title, data in goal_validations:
    add("Stage A", "Goals validation", title,
        "Verify invalid goal confirmation is atomic.", "Baseline completed; current revision known.", data,
        numbered("Record current goals, confirmation, plan, and revision.", "Submit the stated data.", "Reload."),
        "Request is rejected or conflicted; prior goals remain authoritative; no false confirmation or plan appears.",
        "High", "Negative")
add("Stage A", "Goals", "Confirm goals trims values and updates canonical profile",
    "Verify the confirmed goals become durable canonical goals.", "Baseline completed; valid interests saved.",
    "Goals with leading/trailing spaces.",
    numbered("Enter all three goals.", "Confirm.", "Reload programme and career profile."),
    "Whitespace is removed, revision increments once, confirmed timestamp exists, and canonical 3/5-year goals match.",
    "Critical", "Positive", automation="learning-growth.test.ts")
add("Stage A", "Goal integrity", "Later canonical career-goal edit invalidates confirmation",
    "Prevent an outdated learning plan from being presented as confirmed.", "Goals previously confirmed.",
    "Edit canonical three-year goal through the supported profile flow.",
    numbered("Change the canonical career goal.", "Return to Learning & Growth."),
    "Stage returns to goals; confirmation timestamp and generated plan are cleared until reconfirmed.",
    "Critical", "Lifecycle", automation="learning-growth.test.ts")
add("Stage A", "Plan targeting", "Customer-service interest receives support courses",
    "Verify conservative targeting selects relevant content.", "Goals confirmed.",
    "Career area customer_service or exact normalized Customer Support role.",
    numbered("Save the stated interest.", "Confirm goals.", "Inspect plan."),
    "The two support communication courses are planned in the correct order and clearly described as available/planned.",
    "High", "Business rule")
add("Stage A", "Plan targeting", "Software-only interest is not steered to support courses",
    "Prevent irrelevant recommendations.", "Goals confirmed.",
    "Career area software_technology; immediate role Backend developer.",
    numbered("Save the stated interest.", "Inspect plan and refresh."),
    "No support communication course is recommended merely because communication is a learning priority.",
    "High", "Business rule", automation="learning-growth.test.ts")
add("Stage A", "Accessibility", "Keyboard-only completion of Stage A",
    "Ensure the onboarding form is operable without a pointer.", "Candidate at interests.",
    "Keyboard only.",
    numbered("Navigate all fields with Tab/Shift+Tab.", "Select options and submit.", "Complete goals similarly."),
    "Focus order is logical, labels are announced, errors receive focus, and every action is available.",
    "High", "Accessibility")
add("Stage A", "Accessibility", "Screen-reader labels and error association",
    "Ensure fields and stage status are understandable.", "Screen reader enabled.",
    "Invalid and valid Stage A data.",
    numbered("Read the page from the top.", "Trigger validation.", "Review announcements."),
    "Every input has an accessible name; required state, errors, stage, and save result are announced.",
    "High", "Accessibility")

# Stage B: catalog, enrollment, lessons, completion
add("Stage B", "Catalog", "Courses stay locked until baseline and goals are ready",
    "Verify prerequisite enforcement on list and deep links.", "Interests saved; prerequisites incomplete.",
    "Course list and known course URL.",
    numbered("Open course list.", "Open the known course URL directly.", "Attempt enrolment."),
    "Catalog shows locked/empty as designed; deep links and writes cannot bypass prerequisites.",
    "Critical", "Lifecycle", automation="learning-courses.test.ts")
add("Stage B", "Catalog", "Eligible catalog exposes safe course metadata",
    "Confirm catalog correctness without leaking answers.", "Stage A complete.",
    "Support Communication Foundations and In Practice.",
    numbered("Open catalog.", "Open each outline.", "Inspect visible exercise data."),
    "Titles, descriptions, paths, lesson counts, and versions are correct. Correct answers, model answers, and private feedback are not exposed.",
    "Critical", "Security", automation="learning-courses.test.ts")
for path_name, excluded in [("Voice", "chat/email-only"), ("Chat & email", "voice-only")]:
    add("Stage B", "Path selection", f"Enrol in {path_name} path",
        "Verify shared plus selected-path lessons are included.", "Eligible; not enrolled.",
        path_name, numbered("Choose the path.", "Enrol.", "Review lesson list and total."),
        f"Enrollment contains every shared lesson and every {path_name.lower()} lesson, excludes {excluded} lessons, and records the selected path.",
        "Critical", "Positive")
add("Stage B", "Enrollment", "Repeated same-path enrolment is idempotent",
    "Prevent duplicate enrollment rows.", "Candidate already enrolled.",
    "Same course and path.",
    numbered("Submit enrolment again.", "Refresh course and catalog."),
    "The existing enrollment ID and progress are returned; no duplicate or reset occurs.",
    "Critical", "Lifecycle", automation="learning-courses.test.ts")
add("Stage B", "Enrollment", "Cannot switch path after enrolment",
    "Protect the lesson contract for an active enrollment.", "Candidate enrolled in Voice.",
    "Attempt Chat & email enrollment.",
    numbered("Try to enrol in the other path.", "Reload the original course."),
    "Conflict is shown; original path and progress remain unchanged.",
    "High", "Negative")
add("Stage B", "Enrollment", "Concurrent duplicate enrolment",
    "Verify two simultaneous requests create one enrollment.", "Eligible; not enrolled.",
    "Two tabs/actions for same course/path.",
    numbered("Prepare enrolment in two tabs.", "Submit simultaneously.", "Reload catalog."),
    "Both interactions resolve to one enrollment record and one ID; no duplicate rows exist.",
    "Critical", "Concurrency", automation="learning-courses.test.ts")
for bad in ["Unknown course ID", "Unsupported path", "Extra request field", "Forged candidate ID"]:
    add("Stage B", "Enrollment validation", bad,
        "Confirm strict enrolment validation and ownership.", "Eligible candidate.", bad,
        numbered("Record current enrollments.", "Submit the malformed/forged enrolment.", "Reload catalog."),
        "Request is rejected without creating or changing any enrollment.",
        "Critical" if "Forged" in bad else "High", "Security" if "Forged" in bad else "Negative")
add("Stage B", "Lesson draft", "Save a partial lesson draft",
    "Verify unfinished work is private and resumable.", "Active enrollment.",
    "One valid partial free-text answer.",
    numbered("Open a later lesson.", "Enter a partial answer.", "Save draft.", "Leave and return."),
    "Draft persists for the owner, attempts remain zero, and the course resumes at the most recently saved unfinished lesson.",
    "High", "Positive")
add("Stage B", "Lesson access", "Cannot write a lesson from another path",
    "Prevent path bypass.", "Voice enrollment exists.",
    "Chat/email-only lesson ID.",
    numbered("Attempt to open and save the excluded lesson."),
    "Lesson is not found/available; no progress row is created.",
    "Critical", "Security")
add("Stage B", "Lesson validation", "Forged exercise key is rejected",
    "Prevent submission of answers outside the lesson definition.", "Active enrollment.",
    "answers={forged_exercise: value}.",
    numbered("Record lesson progress.", "Submit forged answer.", "Reload lesson."),
    "Request is rejected atomically; attempts, revision, answers, and completion do not change.",
    "Critical", "Security")
add("Stage B", "Lesson validation", "Incomplete submit creates no progress",
    "Ensure required answers are enforced before persistence.", "Untouched lesson.",
    "Empty answers with submit action.",
    numbered("Submit without required answers.", "Reload lesson and inspect progress."),
    "Validation is shown and no draft/progress row is created as a side effect.",
    "High", "Negative", automation="learning-courses.test.ts")
add("Stage B", "Lesson evaluation", "Incorrect choice produces retry feedback",
    "Verify formative feedback without false completion.", "Active lesson with choice exercise.",
    "Select an incorrect option; complete required free text.",
    numbered("Submit the lesson.", "Review status, attempt count, and feedback."),
    "Status remains draft, attempts increments once, incorrect item is identified, and course is not completed.",
    "High", "Positive")
add("Stage B", "Lesson evaluation", "Correct retry completes lesson",
    "Verify correction can complete the lesson.", "Prior failed attempt exists.",
    "Correct option and complete response using current revision.",
    numbered("Correct answers.", "Submit.", "Refresh."),
    "Lesson becomes completed once, completion time exists, and course progress increments once.",
    "Critical", "Positive")
add("Stage B", "Immutability", "Completed lesson cannot be edited",
    "Keep completion evidence stable.", "Lesson completed.",
    "Change a prior answer.",
    numbered("Reopen completed lesson.", "Attempt resubmission or draft save."),
    "Write is rejected; original answers, feedback, attempts, revision, and completion time remain unchanged.",
    "Critical", "Lifecycle")
add("Stage B", "Concurrency", "Two tabs save the same lesson revision",
    "Prevent lost updates.", "Same draft lesson open in two tabs at revision N.",
    "Different answers in each tab.",
    numbered("Save Tab A.", "Save Tab B without refresh.", "Reload both."),
    "One write succeeds and one receives a conflict; the winning answer is not overwritten.",
    "Critical", "Concurrency", automation="learning-courses.test.ts")
add("Stage B", "Versioning", "Stale course version fails closed",
    "Prevent progress against changed course definitions.", "Enrollment references a version not equal to catalog version.",
    "Stale enrollment fixture.",
    numbered("Open course.", "Attempt draft and submit."),
    "A version conflict blocks use; no new progress or reward is written.",
    "Critical", "Compatibility", automation="learning-courses.test.ts")
add("Stage B", "Completion", "Completing all selected lessons completes course",
    "Verify final enrollment invariants.", "All but final selected-path lesson completed.",
    "Valid final answers.",
    numbered("Complete final lesson.", "Return to catalog.", "Refresh."),
    "Enrollment is completed once, completed lesson count equals selected-path total, resume pointer is cleared, and completedAt exists.",
    "Critical", "Lifecycle")
add("Stage B", "Completion", "Concurrent valid final lesson submissions",
    "Prove one completion transition under a race.", "Final lesson ready in two tabs at same revision.",
    "Two identical valid final submissions.",
    numbered("Submit both at nearly the same time.", "Reload course and achievements."),
    "Exactly one succeeds; one conflicts. One course completion and exactly one applicable reward of each type exist.",
    "Critical", "Concurrency", automation="learning-courses.test.ts")
add("Stage B", "Rewards", "First course awards both expected credit events",
    "Verify exact first-course reward accounting.", "Candidate has never completed a course.",
    "Complete first course.",
    numbered("Record balance and ledger.", "Complete final lesson.", "Open achievements."),
    "One +20 completed-course event and one +25 first-course event appear with exact source IDs; balance increases by 45.",
    "Critical", "Business rule")
add("Stage B", "Rewards", "Second course does not repeat first-course reward",
    "Ensure the milestone is awarded once.", "Candidate already completed one course.",
    "Complete another eligible course.",
    numbered("Record ledger.", "Complete second course.", "Open achievements."),
    "One new +20 course event appears; no second +25 first-course event is created.",
    "Critical", "Business rule")
add("Stage B", "Scoring firewall", "Learning does not change hiring scores",
    "Protect the boundary between private development and employer evaluation.", "Candidate has stored hiring/match scores.",
    "Complete lessons and course.",
    numbered("Record all hiring-facing scores.", "Complete course.", "Recheck candidate/employer views."),
    "Hiring, match, résumé-screen, and ranking scores are unchanged by learning participation or results.",
    "Critical", "Privacy")
add("Stage B", "Catalog lifecycle", "Existing enrollment remains accessible after interests change",
    "Avoid stranding in-progress learning.", "Candidate enrolled in a relevant course.",
    "Change current interest to a different career area.",
    numbered("Save new interests.", "Open existing course.", "Inspect new catalog."),
    "Existing private enrollment remains accessible; unrelated new courses are conservatively hidden.",
    "High", "Lifecycle")
add("Stage B", "Usability", "Course resume pointer chooses latest unfinished work",
    "Verify return-to-learning behavior.", "Several lessons have drafts/completions.",
    "Save drafts in non-sequential order.",
    numbered("Save a draft in lesson A.", "Save later draft in lesson B.", "Leave course.", "Resume."),
    "Resume opens the most recently updated unfinished selected-path lesson, never an excluded or completed lesson.",
    "Medium", "Usability")
for condition in ["Offline before save", "Connection drops during save", "Double-click Save", "Refresh while request is pending"]:
    add("Stage B", "Resilience", condition,
        "Ensure network/UI races do not corrupt lesson state.", "Active lesson with unsaved work.", condition,
        numbered("Enter a valid answer.", f"Simulate: {condition}.", "Reconnect/reload and inspect state."),
        "UI gives a clear result; no duplicate attempt/completion/reward occurs; confirmed saved work remains durable.",
        "High", "Resilience")

# Stage C: assessments and reports
add("Stage C", "Assessment start", "Start assessment for selected course path",
    "Verify the task set and rubric are versioned and path-correct.", "Eligible completed/active course as required.",
    "Voice and Chat & email paths.",
    numbered("Start each path assessment with separate candidates.", "Review task set metadata."),
    "Correct course/path/task-set/rubric versions are frozen; tasks belong to the selected path.",
    "Critical", "Positive")
add("Stage C", "Assessment start", "Repeated start returns active draft",
    "Prevent duplicate active assessments.", "Draft assessment exists.",
    "Start same course/path again.",
    numbered("Record assessment ID.", "Start again from another tab.", "Reload."),
    "Existing draft ID and task state are returned; no duplicate draft is created.",
    "Critical", "Lifecycle")
add("Stage C", "Assessment start", "Concurrent start creates one draft",
    "Verify the unique active-attempt rule under race.", "No draft assessment exists.",
    "Two simultaneous Start actions.",
    numbered("Trigger both starts simultaneously.", "Inspect resulting IDs and assessment list."),
    "Both interactions converge on one active draft assessment.",
    "Critical", "Concurrency", automation="learning-assessments.test.ts")
for bad in ["Unknown course", "Unsupported path", "Extra body field", "Forged candidate ID", "Unavailable pilot"]:
    add("Stage C", "Assessment access", bad,
        "Confirm assessment creation is strict and owner-only.", "Known initial assessment count.", bad,
        numbered("Submit the invalid start.", "Reload assessment list."),
        "Request is rejected without creating an assessment or task rows.",
        "Critical" if "Forged" in bad else "High", "Security" if "Forged" in bad else "Negative")
add("Stage C", "Task drafts", "Save partial task response",
    "Verify private draft persistence.", "Draft assessment exists.",
    "Response above zero and below maximum length.",
    numbered("Enter a partial answer.", "Save draft.", "Leave and reopen."),
    "Response, draft status, and revision persist only for the owner; assessment remains incomplete.",
    "High", "Positive")
assessment_validations = [
    ("Blank submitted response", "Submit empty/whitespace response."),
    ("Response below minimum evidence", "Submit fewer non-space characters than required."),
    ("Response exceeds maximum", "Submit more than 6000 characters."),
    ("Unknown task key", "Use a task key outside the frozen task set."),
    ("Extra request field", "Add an unsupported field."),
    ("Fractional revision", "Use revision 0.5."),
]
for title, data in assessment_validations:
    add("Stage C", "Task validation", title,
        "Verify invalid assessment writes are atomic.", "Draft assessment and task state recorded.", data,
        numbered("Record task and parent state.", "Submit stated invalid data.", "Reload."),
        "Validation rejects the write; response, revision, status, parent state, and reward ledger remain unchanged.",
        "High", "Negative")
add("Stage C", "Concurrency", "Two tabs save same assessment task revision",
    "Prevent lost assessment evidence.", "Same task open in two tabs at revision N.",
    "Different valid responses.",
    numbered("Save or submit Tab A.", "Submit Tab B without reload.", "Reload."),
    "One succeeds and one conflicts. The accepted response is not overwritten.",
    "Critical", "Concurrency")
add("Stage C", "Task lifecycle", "Submitted task is immutable",
    "Keep evidence and reports reproducible.", "One task submitted.",
    "Attempt changed response.",
    numbered("Reopen submitted task.", "Attempt draft save and submit."),
    "Writes are rejected; original response, status, revision, and timestamps remain unchanged.",
    "Critical", "Lifecycle")
add("Stage C", "Completion gate", "Assessment remains draft with one incomplete task",
    "Ensure reports are not generated from partial evidence.", "All but one task submitted.",
    "Leave final task draft/not started.",
    numbered("Refresh assessment.", "Attempt completion/report access."),
    "Parent remains draft, no report snapshot/completion time/reward exists, and missing task is identified.",
    "Critical", "Lifecycle")
add("Stage C", "Completion", "Final task creates frozen report",
    "Verify one deterministic completion transition.", "All but final task submitted.",
    "Valid final response.",
    numbered("Record current state.", "Submit final task.", "Reload report repeatedly."),
    "Assessment completes once; report snapshot, rubric/task versions, dimensions, evidence, and recommendations remain byte-for-byte stable.",
    "Critical", "Lifecycle")
add("Stage C", "Concurrency", "Concurrent valid final task submissions",
    "Prove exact-once completion and reward under race.", "Final task open in two tabs at same revision.",
    "Two identical valid submissions.",
    numbered("Submit simultaneously.", "Reload report and achievements."),
    "One succeeds and one conflicts. Exactly one completed parent, frozen report, and +25 first-assessment reward exist.",
    "Critical", "Concurrency", automation="learning-assessments.test.ts")
add("Stage C", "Rewards", "First completed assessment awards 25 credits once",
    "Verify exact assessment reward.", "No prior completed assessment reward.",
    "Complete assessment.",
    numbered("Record balance and ledger.", "Complete final task.", "Retry final submission.", "Reload achievements."),
    "Exactly one +25 first-assessment event with the correct source ID exists; retries do not change balance.",
    "Critical", "Business rule")
add("Stage C", "Repeat attempts", "New assessment can start after completion",
    "Allow reassessment without altering history.", "One assessment completed.",
    "Start same course/path again.",
    numbered("Start a new attempt.", "Compare IDs and prior report."),
    "A fresh draft with a new ID opens. Prior completed assessment/report remains immutable and visible as history where designed.",
    "High", "Lifecycle")
add("Stage C", "Evaluator", "Empty evidence is not observed",
    "Confirm the report does not infer skill from absent evidence.", "Assessment definition available.",
    "Empty/near-empty answer.",
    numbered("Submit through approved minimum/validation fixture or evaluator test screen.", "Inspect report."),
    "Affected dimensions are not_observed; no percentage or hiring score is produced.",
    "Critical", "Business rule")
add("Stage C", "Evaluator", "Unsafe guarantee is downgraded",
    "Ensure unsafe customer promises do not receive strong evidence.", "Assessment task accepts response.",
    "“I guarantee a refund immediately…”",
    numbered("Submit the response in the relevant task.", "Complete assessment.", "Inspect dimension evidence."),
    "Relevant dimension is emerging/developing, not consistent; recommendation points to applicable lessons.",
    "Critical", "Business rule")
add("Stage C", "Evaluator", "Bounded factual response receives consistent evidence",
    "Verify strong evidence is recognised.", "Assessment task accepts response.",
    "Acknowledgement, verified facts, specific next update, no unsafe promise.",
    numbered("Submit equivalent high-quality evidence.", "Complete assessment.", "Inspect report."),
    "Covered dimension can be consistent with evidence grounded in the actual response.",
    "High", "Business rule")
add("Stage C", "Evaluator", "Equivalent response is deterministic",
    "Ensure repeated evaluation cannot drift.", "Same frozen definition and responses.",
    "Identical responses submitted/evaluated repeatedly in a disposable environment.",
    numbered("Evaluate or complete with the same inputs twice.", "Compare reports."),
    "Dimensions, evidence, recommendations, and versions are identical.",
    "Critical", "Compatibility", automation="learning-assessments.test.ts")
add("Stage C", "Report privacy", "Employer cannot access developmental report",
    "Protect candidate-private coaching evidence.", "Completed assessment exists.",
    "Recruiter, tenant admin, and another candidate.",
    numbered("Attempt report URL/API as each non-owner.", "Check existence-leak wording."),
    "Access is denied or not found without exposing report contents or existence.",
    "Critical", "Security")
add("Stage C", "Report content", "Recommendations stay inside selected path",
    "Prevent unusable recommendations.", "Completed assessment in each path.",
    "Voice and Chat & email.",
    numbered("Open each report.", "Check every recommended lesson against course outline."),
    "Every recommendation belongs to the selected path or shared lesson set; no excluded-path lesson appears.",
    "High", "Business rule")
add("Stage C", "Report content", "No percentages or hiring scores appear",
    "Preserve developmental framing.", "Completed assessment.",
    "Report UI and downloaded/printed view if available.",
    numbered("Inspect all headings, cards, network payloads, and print output."),
    "Report uses developmental levels only; no percentage, ranking, employer score, red failure badge, or hiring recommendation appears.",
    "Critical", "Privacy")
for viewport in ["Mobile portrait", "Desktop"]:
    add("Stage C", "Report usability", f"Assessment report is usable on {viewport}",
        "Verify evidence and recommendations remain readable.", "Completed assessment.", viewport,
        numbered("Open report at target viewport.", "Expand all sections.", "Navigate recommendations."),
        "No overflow or obscured evidence; expansion, focus order, and lesson links work.",
        "Medium", "Usability", platform=viewport)

# Stage D: voice cycles, chronology, reviews, comparison
add("Stage D", "Cycle start", "Start first voice-growth cycle",
    "Verify a fresh private cycle uses the correct paired forms.", "Eligible candidate; no active cycle.",
    "Current voice course/version.",
    numbered("Start Stage D.", "Inspect baseline form and cycle metadata."),
    "One cycle starts in baseline state with Form A, correct course/version, no fabricated completion, and no reward.",
    "Critical", "Positive")
add("Stage D", "Cycle start", "Repeated start returns active cycle",
    "Prevent duplicate active cycles.", "Active incomplete cycle exists.",
    "Start action from another tab.",
    numbered("Record cycle ID.", "Start again.", "Reload."),
    "Existing cycle is returned; no duplicate cycle, turns, or reward is created.",
    "Critical", "Lifecycle")
add("Stage D", "Baseline", "Baseline alone gives no credits",
    "Preserve the first-completed-cycle reward rule.", "Cycle in baseline.",
    "Complete Form A only.",
    numbered("Record reward balance.", "Complete baseline.", "Open achievements."),
    "Cycle moves to training, but balance and ledger do not receive a voice-growth reward.",
    "Critical", "Business rule")
add("Stage D", "Chronology", "Form B is blocked before tracked training",
    "Ensure reassessment follows real training evidence.", "Baseline completed; training incomplete.",
    "Direct reassessment URL/action.",
    numbered("Attempt to open/start/submit Form B.", "Reload cycle."),
    "Reassessment is blocked; no Form B answers, report, completion, or reward are created.",
    "Critical", "Lifecycle")
add("Stage D", "Training review", "Review starts idempotently",
    "Verify one tracked review attempt per cycle.", "Cycle in training.",
    "Start review twice, including rapidly.",
    numbered("Start tracked review.", "Start again immediately/from another tab."),
    "Both resolve to one candidate- and tenant-owned attempt; lesson progress is preserved.",
    "Critical", "Concurrency")
add("Stage D", "Training review", "Review lesson list matches exact voice path",
    "Ensure the training contract contains all and only applicable lessons.", "Tracked review started.",
    "Voice course catalog.",
    numbered("Count and list review lessons.", "Compare with shared + voice catalog IDs."),
    "Required catalog lesson IDs match exactly; chat/email-only lessons are excluded.",
    "Critical", "Business rule")
add("Stage D", "Training review", "Wrong lesson cannot satisfy exact-set completion",
    "Prevent count-only completion bypass.", "Review attempt in progress.",
    "Replace one required lesson with an unknown ID in a controlled test fixture.",
    numbered("Open review progress.", "Confirm row count equals expected but exact ID is missing.", "Attempt completion."),
    "Attempt remains in progress and no +10 review reward is awarded.",
    "Critical", "Security", automation="learning-voice-progress.test.ts")
add("Stage D", "Training review", "Adding final exact lesson completes review once",
    "Verify exact-set completion and reward.", "All exact lessons but one reviewed; adversarial extra row may exist.",
    "Review final required lesson.",
    numbered("Mark final required lesson reviewed.", "Repeat the same action.", "Open achievements."),
    "Attempt completes once; cycle records tracked-review evidence; exactly one +10 review reward exists; replay does not alter state.",
    "Critical", "Lifecycle")
add("Stage D", "Training review", "Cannot review another candidate's cycle",
    "Enforce owner privacy.", "Candidate A has active review.",
    "Candidate B in same tenant and Candidate C in another tenant.",
    numbered("Attempt list/start/mark lesson using Candidate A cycle ID as each other candidate."),
    "Access is denied/not found; Candidate A review, cycle, and rewards remain unchanged.",
    "Critical", "Security")
add("Stage D", "Form B", "Tracked review unlocks reassessment",
    "Verify the chronology gate opens only after completion.", "Baseline and tracked review completed.",
    "Return to voice progress.",
    numbered("Refresh Stage D.", "Open reassessment."),
    "State advances to reassessment and Form B is available with fresh prompts.",
    "Critical", "Lifecycle")
add("Stage D", "Paired forms", "Form A and Form B measure same constructs with fresh prompts",
    "Confirm like-for-like comparison without answer memorization.", "Both definitions available.",
    "Form A and Form B.",
    numbered("Compare construct labels, task versions, and prompts.", "Complete both."),
    "Constructs are comparable and version-compatible, while prompts are meaningfully different.",
    "Critical", "Business rule")
add("Stage D", "Form B", "Concurrent valid final responses complete once",
    "Prove exact-once cycle completion under race.", "Fresh candidate; final Form B response open twice.",
    "Two valid submissions at same revision.",
    numbered("Confirm no prior voice reward.", "Submit both simultaneously.", "Reload report and achievements."),
    "One succeeds and one conflicts. One cycle completion, one report, and one exact +50 first-cycle reward exist.",
    "Critical", "Concurrency", automation="learning-voice-progress.test.ts")
add("Stage D", "Rewards", "Only first completed voice-growth cycle awards 50 credits",
    "Prevent recurring credits for repeat practice.", "Candidate completed first cycle and starts another.",
    "Complete second comparable cycle.",
    numbered("Record ledger.", "Complete the second cycle.", "Open achievements."),
    "Second report is saved, but no second +50 first-cycle reward is added.",
    "Critical", "Business rule")
add("Stage D", "Comparison", "Semantically equivalent paired answers preserve levels",
    "Avoid false progress/regression caused by prompt wording.", "Compatible paired forms.",
    "Equivalent quality and meaning in Form A and Form B.",
    numbered("Complete both with equivalent evidence.", "Open comparison report."),
    "Comparable dimensions remain at the same level; report does not invent improvement.",
    "High", "Business rule")
add("Stage D", "Comparison", "Real improvement is supported by like-for-like evidence",
    "Verify progress claims are grounded.", "Compatible baseline and reassessment.",
    "Baseline weak; reassessment specific, safe, and action-oriented.",
    numbered("Complete baseline.", "Finish tracked training.", "Complete reassessment.", "Inspect report."),
    "Improvement is claimed only for comparable dimensions with supporting before/after evidence.",
    "Critical", "Business rule")
add("Stage D", "Comparison", "Incompatible versions fail closed",
    "Prevent invalid longitudinal claims.", "Cycle/report version mismatch fixture.",
    "Incompatible rubric, task-set, or course versions.",
    numbered("Attempt report generation/display.", "Inspect comparison state."),
    "Report says not comparable or blocks comparison; it does not fabricate progress.",
    "Critical", "Compatibility")
add("Stage D", "Comparison", "Unsafe promises cannot score consistent",
    "Protect coaching quality.", "Voice task accepts response.",
    "Guarantee/refund promise without authority or verification.",
    numbered("Use unsafe response in baseline or reassessment.", "Inspect dimension."),
    "Relevant dimension cannot be consistent; feedback encourages bounded, verified commitments.",
    "Critical", "Business rule")
add("Stage D", "Privacy", "Voice-development report stays candidate-private",
    "Prevent developmental evidence entering employer workflows.", "Completed cycle.",
    "Recruiter, admin, other candidate, candidate owner.",
    numbered("Open report as owner.", "Attempt same URL/API as each non-owner."),
    "Owner sees report; all non-owners are denied without an existence leak.",
    "Critical", "Security")
add("Stage D", "Scoring firewall", "Voice progress does not change hiring signals",
    "Separate learning from selection.", "Candidate has employer-facing scores.",
    "Complete baseline, review, and Form B.",
    numbered("Record employer-visible scores and ranking.", "Complete Stage D.", "Recheck employer views."),
    "No hiring score, rank, recommendation, or pipeline stage changes because of private voice practice.",
    "Critical", "Privacy")
for condition in ["Microphone permission denied", "Microphone unavailable", "Audio upload interruption",
                  "Refresh during recording", "Very short/empty recording", "Unsupported audio format"]:
    add("Stage D", "Voice resilience", condition,
        "Verify recoverable failures do not corrupt the cycle.", "Active voice task.", condition,
        numbered("Trigger the stated condition.", "Follow recovery guidance.", "Reload cycle."),
        "Clear actionable error is shown; no false submitted turn/completion/reward exists; prior valid progress remains.",
        "High", "Resilience")
for viewport in ["Mobile portrait", "Desktop"]:
    add("Stage D", "Voice usability", f"Voice cycle is usable on {viewport}",
        "Verify controls and progress remain accessible.", "Active cycle.", viewport,
        numbered("Complete one voice turn.", "Open training review.", "Inspect comparison report."),
        "Recording controls, timers, transcript/status, lesson list, and report are visible and keyboard/screen-reader usable.",
        "High", "Usability", platform=viewport)

# Rewards and achievements
add("Rewards", "Achievements access", "Achievements are candidate-owner only",
    "Verify private reward history cannot be inspected.", "Candidate A has rewards.",
    "Candidate B same tenant; Candidate C other tenant; staff role.",
    numbered("Open achievements as Candidate A.", "Try Candidate A URL/API as each non-owner."),
    "Only Candidate A sees the data; others receive denied/not-found without balances or badge names.",
    "Critical", "Security")
add("Rewards", "Balance", "Credit balance equals immutable ledger sum",
    "Verify exact accounting.", "Candidate has several positive reward events.",
    "Course, first course, assessment, review, and voice-cycle events.",
    numbered("List ledger credits.", "Calculate sum independently.", "Compare with displayed balance."),
    "Displayed balance equals the exact ledger sum; no hidden or duplicated credits exist.",
    "Critical", "Business rule")
reward_matrix = [
    ("First completed course", 25, "first_completed_course"),
    ("Completed course", 20, "completed_course"),
    ("First completed assessment", 25, "first_completed_assessment"),
    ("Tracked course review", 10, "completed_course_review"),
    ("First completed voice-growth cycle", 50, "completed_voice_growth"),
]
for title, credits, event in reward_matrix:
    add("Rewards", "Ledger", f"{title} records exact +{credits} event",
        "Verify title, type, source, and credit value.", "Candidate is one valid action away from event.",
        event, numbered("Record ledger.", f"Complete action: {title}.", "Refresh achievements."),
        f"One new {event} row appears with +{credits}, correct source type/ID, candidate, tenant, and earned time.",
        "Critical", "Business rule")
add("Rewards", "Idempotency", "Replaying completed actions never duplicates credits",
    "Verify reward event keys are exact-once.", "All reward event types exist.",
    "Refresh, back/forward, repeated PUT/POST, and two-tab replay.",
    numbered("Record balance and ledger.", "Replay every completed action.", "Reload achievements."),
    "Balance and ledger are unchanged; every event key remains unique.",
    "Critical", "Concurrency")
add("Rewards", "Badges", "Duplicate badge events display one badge",
    "Prevent repeated badge cards while retaining ledger history.", "Two ledger events share one badge key.",
    "Same badge key with different event times.",
    numbered("Open achievements.", "Count badge cards and inspect earned date."),
    "One badge is displayed for the key, using the earliest earned event according to the product rule.",
    "High", "Business rule")
add("Rewards", "Ordering", "Badges have deterministic equal-time ordering",
    "Prevent cards jumping between reloads.", "Multiple badges share earnedAt.",
    "Equal timestamps and stable distinct IDs.",
    numbered("Reload achievements repeatedly.", "Compare order across sessions."),
    "Order is identical on every load using the documented timestamp plus stable secondary ordering.",
    "High", "Compatibility")
add("Rewards", "Recent activity", "Activity is newest-first and capped at 20",
    "Keep the feed stable and bounded.", "Candidate has at least 25 ledger events.",
    "Include equal timestamps.",
    numbered("Open achievements.", "Count and order activity.", "Reload several times."),
    "Exactly 20 newest events are shown, newest-first, with deterministic tie ordering.",
    "High", "Business rule")
add("Rewards", "Milestones", "Milestones derive from completed source records",
    "Verify milestone progress cannot be inflated by duplicate badge cards.", "Dedicated completed course/assessment/review/cycle fixtures.",
    "Known source counts.",
    numbered("Open achievements.", "Compare milestone counts with source records."),
    "Each milestone reflects the correct completed source records and never counts another candidate's work.",
    "Critical", "Business rule")
add("Rewards", "Privacy", "Achievements response is private and no-store",
    "Prevent cached reward exposure.", "Candidate has rewards.",
    "Inspect response headers.",
    numbered("Open achievements.", "Inspect Cache-Control.", "Switch accounts in same browser."),
    "Response is private, no-store and account switching never shows previous candidate data.",
    "Critical", "Security")
add("Rewards", "No rewards", "New candidate sees a valid zero state",
    "Verify empty state is honest and actionable.", "Eligible candidate with no ledger rows.",
    "Zero rewards.",
    numbered("Open achievements.", "Inspect balance, badges, milestones, and activity."),
    "Balance is 0; no fabricated badge/activity appears; next eligible learning action is clear.",
    "Medium", "Usability")
add("Rewards", "Scoring firewall", "Credits and badges never appear as hiring merit",
    "Protect candidate developmental privacy.", "Candidate has high learning balance.",
    "Recruiter candidate list, profile, matching, evaluation, and export surfaces.",
    numbered("Record employer views before rewards.", "Earn rewards.", "Recheck every employer surface."),
    "Learning credits, badges, assessment levels, and private progress are absent and do not affect hiring scores.",
    "Critical", "Privacy")

# Cross-cutting security, data isolation, cleanup, and operations
for resource in ["Learning profile", "Course enrollment", "Lesson progress", "Assessment",
                 "Assessment task", "Voice cycle", "Voice turn", "Training review",
                 "Reward ledger"]:
    add("Privacy & Security", "Database RLS", f"{resource} requires tenant and candidate context",
        "Verify defense-in-depth isolation under the application DB role.",
        "Disposable owner row; lexy_app role available in controlled QA environment.",
        "Owner context, same-tenant other candidate, foreign tenant, and missing candidate context.",
        numbered("Query as correct tenant + candidate.", "Repeat as same-tenant other candidate.",
                 "Repeat as foreign tenant.", "Repeat with candidate context unset."),
        "Only the exact owner context can read the row. Every wrong or missing context sees zero rows.",
        "Critical", "Security", "Database tester", "Database",
        "0064/0065/0066 RLS migrations")
add("Privacy & Security", "Ownership", "Forged candidate IDs are ignored or rejected",
    "Ensure identity always comes from authenticated session.", "Candidate A and B exist.",
    "Send Candidate B ID in Candidate A request body/query/path where unsupported.",
    numbered("Capture a valid Candidate A request.", "Inject Candidate B identifier.", "Submit and inspect both accounts."),
    "Request is rejected; no record is read, created, or changed for either candidate.",
    "Critical", "Security")
add("Privacy & Security", "Existence leaks", "Unknown and unauthorized IDs reveal equivalent information",
    "Prevent resource enumeration.", "Known foreign IDs and random nonexistent IDs.",
    "Course, assessment, cycle, review, lesson IDs.",
    numbered("Request each foreign ID.", "Request a random ID.", "Compare status and response detail."),
    "Responses do not reveal which IDs exist or who owns them.",
    "Critical", "Security")
add("Privacy & Security", "Session switching", "Sign-out clears candidate learning state",
    "Prevent client cache leakage.", "Candidate A and B in same browser.",
    "A has progress; B has none.",
    numbered("Open several learning pages as A.", "Sign out.", "Sign in as B."),
    "No A content flashes or remains in cache; B starts from their own state.",
    "Critical", "Security")
add("Privacy & Security", "Data export", "General candidate/employer exports exclude private learning",
    "Keep development data outside hiring and admin exports.", "Candidate has all stages completed.",
    "Available CSV/PDF/API exports.",
    numbered("Export candidate data from recruiter/admin surfaces.", "Search output for learning fields."),
    "Private interests, goals, lessons, assessment evidence, voice reports, rewards, and badges are absent unless a dedicated candidate-owned export is explicitly designed.",
    "Critical", "Privacy")
add("Resilience", "Transaction integrity", "Failure during final write rolls back completion and reward together",
    "Ensure no partial completion or orphan reward.", "Controlled QA fault injection around final transaction.",
    "Course, assessment, review, and voice finalization.",
    numbered("Inject a failure after progress write but before reward/parent completion.", "Submit final action.", "Inspect all related tables."),
    "Transaction rolls back consistently; retry can complete normally; no partial parent, report, or reward exists.",
    "Critical", "Resilience", "QA engineer", "API/Database")
add("Resilience", "Cleanup", "Disposable manual-test data can be removed without orphans",
    "Verify referential lifecycle and repeatable QA.", "Dedicated QA tenant/candidates with all stages.",
    "Delete via approved cleanup process in non-production.",
    numbered("Count all learning child rows.", "Run approved QA cleanup.", "Check related tables."),
    "All disposable rows are removed in safe child-to-parent order or by validated cascades; unrelated candidates remain.",
    "High", "Lifecycle", "QA engineer", "Database")
add("Resilience", "Auditability", "Completion and reward timestamps are internally consistent",
    "Make investigation possible.", "Candidate completes multiple stages.",
    "Server times and ledger/source timestamps.",
    numbered("Record source completion times.", "Compare reward earned times and report snapshots."),
    "Reward cannot precede qualifying completion; frozen report/version metadata remains attributable to the source event.",
    "High", "Lifecycle", "QA engineer", "API/Database")


def build_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    navy, blue, cyan = "0B1F33", "1663A6", "4DA3D9"
    pale, white, grid = "EAF3F8", "FFFFFF", "C7D5E0"
    header_fill = PatternFill("solid", fgColor=navy)
    sub_fill = PatternFill("solid", fgColor=blue)
    thin = Side(style="thin", color=grid)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    readme = wb.create_sheet("Read Me")
    readme.sheet_view.showGridLines = False
    readme["A1"] = "Lexy Learning & Growth — Manual QA Workbook"
    readme["A1"].font = Font(size=20, bold=True, color=white)
    readme["A1"].fill = header_fill
    readme.merge_cells("A1:H2")
    readme["A4"] = "Purpose"
    readme["B4"] = "Manual verification for Stages A–D, courses, assessments, voice growth, achievements, privacy, concurrency, validation, and failure handling."
    readme["A6"] = "Recommended execution"
    readme["B6"] = "Run Critical cases first, then High, then Medium. Use disposable QA candidates and never manipulate production records directly."
    readme["A8"] = "Status values"
    readme["B8"] = "Not Run, Pass, Fail, Blocked, Not Applicable"
    readme["A10"] = "Evidence"
    readme["B10"] = "Record actual result, tester, date, defect ID, and concise notes. Attach screenshots/network evidence in your team’s defect tracker."
    readme["A12"] = "Important privacy rule"
    readme["B12"] = "Learning and Growth is candidate-private. Recruiter/admin access, cached cross-account data, or changes to hiring scores are release-blocking failures."
    readme["A14"] = "Workbook generated"
    readme["B14"] = date.today().isoformat()
    for row in range(4, 15, 2):
        readme[f"A{row}"].font = Font(bold=True, color=blue)
        readme[f"A{row}"].alignment = Alignment(vertical="top")
        readme[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 25
    readme.column_dimensions["B"].width = 115

    test_data = wb.create_sheet("Test Data Matrix")
    test_data.append(["Persona / Fixture", "Purpose", "Required State", "Restrictions"])
    fixture_rows = [
        ("Candidate New", "Stage A and zero-state checks", "Pilot eligible; no learning profile", "Disposable account"),
        ("Candidate Same Tenant", "Same-tenant privacy tests", "Different user/candidate; no shared records", "Never reuse owner session"),
        ("Candidate Foreign Tenant", "Cross-tenant isolation", "Separate QA tenant", "No shared hierarchy"),
        ("Candidate Stage B", "Course/lesson tests", "Stage A complete; not yet enrolled", "Reset between path variants"),
        ("Candidate Stage C", "Assessment tests", "Eligible course/path; no active assessment", "Use separate candidate for concurrency"),
        ("Candidate Stage D", "Voice cycle tests", "Eligible voice course/version", "Use test audio only"),
        ("Candidate Rewards", "Achievements accounting", "Dedicated seeded/completed source events", "Do not depend on another case"),
        ("Recruiter", "Negative access tests", "Same tenant as an owner candidate", "Must never see private learning"),
        ("Tenant Admin", "Negative access tests", "Same tenant as an owner candidate", "Must never see private learning"),
        ("Database QA User", "RLS verification", "Controlled non-production access as lexy_app", "No admin/bypass role for assertions"),
    ]
    for row in fixture_rows:
        test_data.append(row)

    master = wb.create_sheet("All Test Cases")
    master.append(HEADERS)
    for case in cases:
        master.append([case[h] for h in HEADERS])

    stage_order = ["Programme", "Stage A", "Stage B", "Stage C", "Stage D", "Rewards", "Privacy & Security", "Resilience"]
    for stage in stage_order:
        title = stage.replace("Privacy & Security", "Privacy").replace("Programme", "Programme")
        ws = wb.create_sheet(title[:31])
        ws.append(HEADERS)
        for case in cases:
            if case["Stage"] == stage:
                ws.append([case[h] for h in HEADERS])

    defects = wb.create_sheet("Defect Log")
    defect_headers = ["Defect ID", "Test Case ID", "Summary", "Severity", "Status", "Owner", "Date Raised", "Environment", "Evidence / Link", "Resolution Notes"]
    defects.append(defect_headers)
    for _ in range(30):
        defects.append(["", "", "", "", "Open", "", "", "", "", ""])

    summary = wb.create_sheet("Execution Summary", 1)
    summary.sheet_view.showGridLines = False
    summary["A1"] = "Execution Summary"
    summary["A1"].font = Font(size=20, bold=True, color=white)
    summary["A1"].fill = header_fill
    summary.merge_cells("A1:H2")
    summary.append([])
    summary.append(["Stage", "Total", "Not Run", "Pass", "Fail", "Blocked", "N/A", "Pass Rate"])
    for idx, stage in enumerate(stage_order, 5):
        summary.cell(idx, 1, stage)
        summary.cell(idx, 2, f'=COUNTIF(\'All Test Cases\'!$B:$B,A{idx})')
        summary.cell(idx, 3, f'=COUNTIFS(\'All Test Cases\'!$B:$B,A{idx},\'All Test Cases\'!$O:$O,"Not Run")')
        summary.cell(idx, 4, f'=COUNTIFS(\'All Test Cases\'!$B:$B,A{idx},\'All Test Cases\'!$O:$O,"Pass")')
        summary.cell(idx, 5, f'=COUNTIFS(\'All Test Cases\'!$B:$B,A{idx},\'All Test Cases\'!$O:$O,"Fail")')
        summary.cell(idx, 6, f'=COUNTIFS(\'All Test Cases\'!$B:$B,A{idx},\'All Test Cases\'!$O:$O,"Blocked")')
        summary.cell(idx, 7, f'=COUNTIFS(\'All Test Cases\'!$B:$B,A{idx},\'All Test Cases\'!$O:$O,"Not Applicable")')
        summary.cell(idx, 8, f'=IFERROR(D{idx}/(B{idx}-C{idx}-G{idx}),0)')
        summary.cell(idx, 8).number_format = "0.0%"
    total_row = 5 + len(stage_order)
    summary.cell(total_row, 1, "TOTAL")
    for col in range(2, 8):
        letter = get_column_letter(col)
        summary.cell(total_row, col, f"=SUM({letter}5:{letter}{total_row-1})")
    summary.cell(total_row, 8, f'=IFERROR(D{total_row}/(B{total_row}-C{total_row}-G{total_row}),0)')
    summary.cell(total_row, 8).number_format = "0.0%"
    summary["A16"] = "Priority"
    summary["B16"] = "Count"
    for i, p in enumerate(["Critical", "High", "Medium", "Low"], 17):
        summary.cell(i, 1, p)
        summary.cell(i, 2, f'=COUNTIF(\'All Test Cases\'!$J:$J,A{i})')

    chart = BarChart()
    chart.title = "Test Cases by Stage"
    chart.y_axis.title = "Cases"
    chart.x_axis.title = "Stage"
    chart.add_data(Reference(summary, min_col=2, min_row=4, max_row=12), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=5, max_row=12))
    chart.height, chart.width = 8, 15
    summary.add_chart(chart, "J4")

    status_dv = DataValidation(type="list", formula1='"Not Run,Pass,Fail,Blocked,Not Applicable"', allow_blank=False)
    priority_dv = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=False)
    defect_status_dv = DataValidation(type="list", formula1='"Open,In Progress,Ready to Retest,Closed,Rejected"', allow_blank=True)
    defect_sev_dv = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=True)

    case_sheet_names = ["All Test Cases"] + [s.replace("Privacy & Security", "Privacy")[:31] for s in stage_order]
    widths = [15, 18, 24, 38, 38, 42, 35, 58, 60, 12, 16, 17, 22, 28, 16, 45, 18, 16, 16, 35]
    for name in case_sheet_names:
        ws = wb[name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True, color=white)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[1].height = 36
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
            row[0].font = Font(bold=True, color=blue)
        max_row = max(ws.max_row, 2)
        status_dv_local = DataValidation(type="list", formula1='"Not Run,Pass,Fail,Blocked,Not Applicable"', allow_blank=False)
        priority_dv_local = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=False)
        ws.add_data_validation(status_dv_local)
        ws.add_data_validation(priority_dv_local)
        status_dv_local.add(f"O2:O{max_row}")
        priority_dv_local.add(f"J2:J{max_row}")
        ws.conditional_formatting.add(f"O2:O{max_row}", FormulaRule(formula=['$O2="Pass"'], fill=PatternFill("solid", fgColor="C6EFCE")))
        ws.conditional_formatting.add(f"O2:O{max_row}", FormulaRule(formula=['$O2="Fail"'], fill=PatternFill("solid", fgColor="FFC7CE")))
        ws.conditional_formatting.add(f"O2:O{max_row}", FormulaRule(formula=['$O2="Blocked"'], fill=PatternFill("solid", fgColor="FFEB9C")))
        table = Table(displayName=("Cases" + "".join(ch for ch in name if ch.isalnum()))[:250], ref=f"A1:T{max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        ws.add_table(table)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"

    for ws in [test_data, defects]:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True, color=white)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = border
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 28 if col != 3 else 45

    defects.add_data_validation(defect_status_dv)
    defects.add_data_validation(defect_sev_dv)
    defect_status_dv.add("E2:E200")
    defect_sev_dv.add("D2:D200")

    for cell in summary[4]:
        cell.fill = sub_fill
        cell.font = Font(bold=True, color=white)
        cell.alignment = Alignment(horizontal="center")
    for row in summary.iter_rows(min_row=5, max_row=total_row, min_col=1, max_col=8):
        for cell in row:
            cell.border = border
    for col, width in zip(range(1, 9), [24, 12, 12, 12, 12, 12, 12, 14]):
        summary.column_dimensions[get_column_letter(col)].width = width

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


if __name__ == "__main__":
    build_workbook()
    # Read-back validation catches corrupt archives and verifies core workbook features.
    check = load_workbook(OUTPUT, data_only=False)
    assert "All Test Cases" in check.sheetnames
    assert "Execution Summary" in check.sheetnames
    assert check["All Test Cases"].max_row == len(cases) + 1
    assert check["All Test Cases"]["A2"].value
    print(f"Created {OUTPUT} with {len(cases)} manual test cases across {len(check.sheetnames)} sheets.")